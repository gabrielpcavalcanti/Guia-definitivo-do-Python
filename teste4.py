import os
import openpyxl
from tkinter import filedialog

def choose_file():

    try:
        path_file = filedialog.askopenfilename()
        
        return path_file
    
    except OSError:
        print("Select a file")


def creating_list_dados():
    for row in ws.iter_rows(values_only=True):
    
        lista_dados_temporaria = list(row)
        lista_dados.append(lista_dados_temporaria)
    
    return lista_dados

def creating_list_valores():

    lista_valores = list(range(16,len(lista_dados), 25))
    return lista_valores


def creating_database(lista_dados, lista_valores):

    for row_lista in lista_dados:

        if lista_dados.index(row_lista) in lista_valores:
            ws1.append(row_lista)


def main():

    global lista_dados
    lista_dados = []
    global lista_valores
    lista_valores = []
    
    global wb
    wb = openpyxl.load_workbook(choose_file())

    global ws 
    ws = wb.active

    global ws1
    wb.create_sheet("Organized Data", 1)
    ws1 = wb['Organized Data']

    lista_dados = creating_list_dados()
    lista_valores = creating_list_valores()

    ws1.append(lista_dados[14])
    
    creating_database(lista_dados, lista_valores)

    wb.save(choose_file())

main()